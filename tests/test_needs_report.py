"""Тесты отчёта-формы для аптеки (report.py): агрегация по разделу, .xlsx, HTML.

Фикстура: временный YAML-каталог (2 базы × 2 точки; группы Неспецифика,
Медикаменты, Растворы) + SQLite-БД на tmp_path + NeedsService, чтобы
сохранять заявки со снимками unit/grp как в проде (ТЗ: отчёты по снимкам).
Заявки теперь раздельные по разделу (category): растворы-строки — в
solutions-заявку, остальные — в medicaments-заявку.
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from docapp.domain.employee import NURSE
from docapp.needs.catalog import CATEGORY_MEDICAMENTS, CATEGORY_SOLUTIONS, Catalog
from docapp.needs.report import aggregate_requests, build_xlsx, html_table, week_label
from docapp.needs.service import NeedsService
from docapp.needs.store import SqliteNeedsStore

CATALOG_YAML = """\
bases:
  Ленская: [травма, урология]
  Таймырская: [экстренная, гной]

groups:
  Неспецифика:
    Атропин: амп
    Дексаметазон: амп
  Медикаменты:
    Адреналин: амп
  Растворы:
    Физ 200/250: фл
    Рингер: фл
"""

BASE = "Ленская"
WEEK = "2026-08-10"

SOL = CATEGORY_SOLUTIONS
MED = CATEGORY_MEDICAMENTS


@pytest.fixture
def env(tmp_path):
    """Каталог + хранилище + сервис на временных файлах (2 базы × 2 точки)."""
    catalog_path = tmp_path / "catalog.yaml"
    catalog_path.write_text(CATALOG_YAML, encoding="utf-8")
    catalog = Catalog(catalog_path)
    store = SqliteNeedsStore(tmp_path / "needs.db")
    service = NeedsService(store, catalog)
    yield service, store, catalog
    store.close_conn()


def send(service, base, point, category, lines):
    """Сохранить заявку раздела медсестрой и отправить её (status='sent')."""
    service.save(
        user_id=7,
        role=NURSE,
        base=base,
        point=point,
        category=category,
        week_start=WEEK,
        lines=[{"item": item, "qty": qty} for item, qty in lines],
    )
    service.submit(7, NURSE, base, point, category, WEEK)


def aggregate(store, catalog, base=BASE, week=WEEK, category=None):
    """Собрать отчёт по заявкам базы и раздела за неделю (как это сделает роутер T6)."""
    if category is None:
        requests = store.list_range(week, week, base=base)
    else:
        requests = store.list_requests(base, category, week)
    return aggregate_requests(requests, base, category, week, catalog)


class TestAggregate:
    def test_solutions_per_point_and_groups_summed(self, env):
        """a) Растворы поточково (каждая точка, ИТОГО); остальное суммарно по группам."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10)])
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        send(service, BASE, "урология", MED, [("Атропин", 3), ("Дексаметазон", 2)])
        send(service, BASE, "урология", SOL, [("Рингер", 4)])
        agg = aggregate(store, catalog)

        # растворы — поточково, у отсутствующей точки 0, единица из снимка
        assert agg["solutions"]["Физ 200/250"] == {
            "травма": 5, "урология": 0, "ИТОГО": 5, "unit": "фл",
        }
        assert agg["solutions"]["Рингер"] == {
            "травма": 0, "урология": 4, "ИТОГО": 4, "unit": "фл",
        }
        # остальные группы — суммарно по препарату
        assert agg["groups"]["Неспецифика"]["Атропин"] == {"unit": "амп", "qty": 13}
        assert agg["groups"]["Неспецифика"]["Дексаметазон"] == {"unit": "амп", "qty": 2}
        # группы в порядке каталога; пустая группа (Адреналин не заказан) не выводится
        assert list(agg["groups"]) == ["Неспецифика"]

    def test_solutions_section_only(self, env):
        """a') category='solutions' → только поточковый свод, groups пусто."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10)])
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        agg = aggregate(store, catalog, category=SOL)

        assert agg["solutions"]["Физ 200/250"]["ИТОГО"] == 5
        assert agg["groups"] == {}

    def test_medicaments_section_only(self, env):
        """a'') category='medicaments' → только свод по группам, solutions пусто."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10)])
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        agg = aggregate(store, catalog, category=MED)

        assert agg["solutions"] == {}
        assert agg["groups"]["Неспецифика"]["Атропин"] == {"unit": "амп", "qty": 10}

    def test_zero_qty_and_drafts_excluded(self, env):
        """b) Строки с qty == 0 и черновики (draft) в отчёт не попадают."""
        service, store, catalog = env
        # отправленная: Атропин есть, Дексаметазон с qty == 0
        send(service, BASE, "травма", MED, [("Атропин", 10), ("Дексаметазон", 0)])
        # черновик (не отправлен) с реальными количествами
        service.save(
            8, NURSE, BASE, "урология", MED, WEEK, [{"item": "Атропин", "qty": 4}]
        )
        agg = aggregate(store, catalog)

        assert agg["solutions"] == {}
        assert agg["groups"]["Неспецифика"] == {
            "Атропин": {"unit": "амп", "qty": 10}
        }
        assert "Дексаметазон" not in agg["groups"]["Неспецифика"]

    def test_points_and_zero_at_missing_point(self, env):
        """c) 'points' — все точки базы; у раствора отсутствующая точка = 0."""
        service, store, catalog = env
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        agg = aggregate(store, catalog, category=SOL)

        assert agg["points"] == ["травма", "урология"]
        assert agg["solutions"]["Физ 200/250"]["травма"] == 5
        assert agg["solutions"]["Физ 200/250"]["урология"] == 0
        assert agg["solutions"]["Физ 200/250"]["ИТОГО"] == 5

    def test_item_with_zero_total_not_in_groups(self, env):
        """d) Препарат с суммарным qty == 0 не появляется в 'groups'."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10), ("Адреналин", 0)])
        agg = aggregate(store, catalog, category=MED)

        assert "Медикаменты" not in agg["groups"]
        assert agg["groups"]["Неспецифика"] == {
            "Атропин": {"unit": "амп", "qty": 10}
        }


class TestWeekLabel:
    def test_monday_to_sunday_range(self):
        """e) week_label('2026-08-10') == '10.08.2026 – 16.08.2026' (пн–вс)."""
        assert week_label("2026-08-10") == "10.08.2026 – 16.08.2026"


class TestBuildXlsx:
    def test_two_sheets_with_header_rows_and_totals(self, env):
        """f) Оба раздела: листы «Растворы» и «Медикаменты» с шапкой и подытогом."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10)])
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        send(service, BASE, "урология", MED, [("Атропин", 3)])
        send(service, BASE, "урология", SOL, [("Рингер", 4)])
        agg = aggregate(store, catalog)

        wb = load_workbook(BytesIO(build_xlsx(agg)))
        assert wb.sheetnames == ["Растворы", "Медикаменты"]

        # ── лист «Растворы» ──────────────────────────────────────────
        sol = wb["Растворы"]
        rows = [[cell.value for cell in row] for row in sol.iter_rows()]
        assert any("База: Ленская" == v for row in rows for v in row if v)
        assert any("Неделя: 10.08.2026 – 16.08.2026" == v for row in rows for v in row if v)
        assert any(str(v).startswith("Сформирован: ") for row in rows for v in row if v)
        header = next(row for row in rows if row and row[0] == "Раствор")
        assert header[1] == "Ед."
        assert header[2] == "травма" and header[3] == "урология"
        assert header[-1] == "ИТОГО"
        row_fiz = next(row for row in rows if row and row[0] == "Физ 200/250")
        assert row_fiz[1] == "фл"
        assert row_fiz[2] == 5 and row_fiz[3] == 0 and row_fiz[-1] == 5

        # ── лист «Медикаменты» ───────────────────────────────────────
        need = wb["Медикаменты"]
        nrows = [[cell.value for cell in row] for row in need.iter_rows()]
        assert any(row and row[0] == "Неспецифика" for row in nrows)  # секция группы
        assert any(row and row[0] == "Атропин" and row[2] == 13 for row in nrows)
        assert any(row and row[0] == "Итого по группе: 13" for row in nrows)  # подытог

    def test_solutions_only_sheet(self, env):
        """f') Отчёт одного раздела растворов — только лист «Растворы»."""
        service, store, catalog = env
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        agg = aggregate(store, catalog, category=SOL)

        wb = load_workbook(BytesIO(build_xlsx(agg)))
        assert wb.sheetnames == ["Растворы"]

    def test_medicaments_only_sheet(self, env):
        """f'') Отчёт одного раздела медикаментов — только лист «Медикаменты»."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10)])
        agg = aggregate(store, catalog, category=MED)

        wb = load_workbook(BytesIO(build_xlsx(agg)))
        assert wb.sheetnames == ["Медикаменты"]
        need = wb["Медикаменты"]
        nrows = [[cell.value for cell in row] for row in need.iter_rows()]
        assert any(row and row[0] == "Неспецифика" for row in nrows)


class TestHtmlTable:
    def test_groups_items_and_escaping(self, env):
        """g) Названия групп и препаратов есть; '<script>' не ломает таблицу."""
        service, store, catalog = env
        send(service, BASE, "травма", MED, [("Атропин", 10)])
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        # позиция с HTML-тегом в имени: в каталоге её нет, снимок группы пуст —
        # попадает в секцию «Без группы», имя обязано экранироваться
        service.save(8, NURSE, BASE, "урология", MED, WEEK, [{"item": "<script>", "qty": 2}])
        service.submit(8, NURSE, BASE, "урология", MED, WEEK)
        agg = aggregate(store, catalog)

        html = html_table(agg)
        assert "Неспецифика" in html
        assert "Атропин" in html
        assert "Физ 200/250" in html
        assert "Без группы" in html
        assert "Итого по группе" in html
        assert "report-table" in html  # стилизованная таблица просмотра
        assert "&lt;script&gt;" in html  # имя экранировано
        assert "<script>" not in html  # сырой тег не попал в разметку

    def test_solutions_only_html(self, env):
        """g') Отчёт одного раздела растворов: нет секций групп."""
        service, store, catalog = env
        send(service, BASE, "травма", SOL, [("Физ 200/250", 5)])
        agg = aggregate(store, catalog, category=SOL)

        html = html_table(agg)
        assert "Растворы" in html
        assert "Физ 200/250" in html
        assert "Неспецифика" not in html
        assert "Итого по группе" not in html
