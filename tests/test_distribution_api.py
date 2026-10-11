"""Тесты «Распределения» через HTTP: страница, доступ, данные и запись из соседа.

Раздел целиком принадлежит заведующему (`distribution.manage`, ADR-0024): врач,
медсестра и старшая сестра не видят ни пункта меню, ни данных.

С 11.10.2026 прогон и публикация — два разных шага (ADR-0027): загрузка файла
показывает результат на странице и **не** ставит меток, файл отдаёт отдельная
кнопка, а месяц закрывает «Опубликовать». Здесь проверяется именно это, потому
что прежнее поведение («метки по факту вкидывания файла») владелец отменил.

Ходит раздел за записями к модулю `records`, поэтому здесь же проверяется, что
запись, поданная через главную страницу, сразу видна в распределении — модули
смотрят на одни данные через интерфейсы, а не каждый в свою копию.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from docapp.auth.passwords import hash_password
from docapp.domain.employee import DOCTOR, HEAD, HEAD_NURSE, NURSE, Employee
from docapp.people.store import SqliteEmployeeStore
from docapp.web.app import create_app
from factories import make_db, vedomost_xlsx

#: Тип файла .xlsx — как его присылает браузер.
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PASSWORDS = {
    "head": "headpass",
    "doc": "docpass",
    "nurse": "nursepass",
    "headnurse": "hnpass",
}


def _seed(db_path) -> dict:
    """Заведующий, врач, медсестра и старшая сестра с логинами."""
    with SqliteEmployeeStore(db_path) as employees:
        head = employees.add(
            Employee(last_name="Иванов", first_name="Иван", role=HEAD,
                     login="head", password_hash=hash_password(PASSWORDS["head"]))
        )
        doctor = employees.add(
            Employee(last_name="Петров", first_name="Пётр", role=DOCTOR,
                     login="doc", password_hash=hash_password(PASSWORDS["doc"]))
        )
        nurse = employees.add(
            Employee(last_name="Сидорова", first_name="Анна", role=NURSE,
                     login="nurse", password_hash=hash_password(PASSWORDS["nurse"]))
        )
        head_nurse = employees.add(
            Employee(last_name="Волкова", first_name="Вера", role=HEAD_NURSE,
                     login="headnurse", password_hash=hash_password(PASSWORDS["headnurse"]))
        )
    return {"head": head, "doctor": doctor, "nurse": nurse, "head_nurse": head_nurse}


@pytest.fixture
def client(tmp_path):
    db = make_db(tmp_path, seed=False)
    ids = _seed(db)
    app = create_app(db_path=db, secret="test-secret")
    with TestClient(app, follow_redirects=False) as c:
        c.ids = ids
        yield c


def login(client, who: str) -> None:
    """Войти под сотрудником: head/doc/nurse/headnurse."""
    response = client.post("/login", data={"login": who, "password": PASSWORDS[who]})
    assert response.status_code == 303


class TestPage:
    def test_requires_login(self, client):
        assert client.get("/distribution").status_code == 303
        assert client.get("/distribution").headers["location"] == "/login"

    def test_vedomost_page_for_head(self, client):
        """Главная страница раздела — загрузка ведомости."""
        login(client, "head")
        response = client.get("/distribution")
        assert response.status_code == 200
        assert "Распределение" in response.text
        assert 'enctype="multipart/form-data"' in response.text
        assert 'action="/distribution/spread"' in response.text

    def test_page_describes_the_format_of_the_file(self, client):
        """Формат подаваемого файла написан на странице, а не только в ТЗ (ADR-0027)."""
        login(client, "head")
        text = client.get("/distribution").text

        assert "Обязательные столбцы" in text
        assert "Ф.И.О. пациента" in text
        assert "первый лист" in text

    def test_counts_page_for_head(self, client):
        login(client, "head")
        response = client.get("/distribution/counts")
        assert response.status_code == 200
        assert "Всего записей" in response.text

    def test_menu_has_distribution(self, client):
        login(client, "head")
        assert 'href="/distribution"' in client.get("/").text

    def test_bad_split_is_a_message_not_a_crash(self, client):
        """Неизвестный разрез — страница с ошибкой 400, а не 500."""
        login(client, "head")
        response = client.get("/distribution/counts?by=base")
        assert response.status_code == 400
        assert "разрез" in response.text


class TestAccess:
    """Раздел закрыт для всех, кроме заведующего (ADR-0024)."""

    @pytest.mark.parametrize("who", ["doc", "nurse", "headnurse"])
    def test_page_is_forbidden(self, client, who):
        login(client, who)
        assert client.get("/distribution").status_code == 403
        assert client.get("/distribution/counts").status_code == 403

    @pytest.mark.parametrize("who", ["doc", "nurse", "headnurse"])
    def test_api_is_forbidden(self, client, who):
        login(client, who)
        assert client.get("/distribution/api/aggregate").status_code == 403

    @pytest.mark.parametrize("who", ["doc", "nurse", "headnurse"])
    def test_upload_is_forbidden(self, client, who):
        login(client, who)
        response = client.post(
            "/distribution/spread",
            data={"month": "2026-06"},
            files={"file": ("ведомость.xlsx", "не файл".encode("utf-8"), "application/octet-stream")},
        )
        assert response.status_code == 403

    @pytest.mark.parametrize("who", ["doc", "nurse", "headnurse"])
    def test_download_is_forbidden(self, client, who):
        login(client, who)
        assert client.get("/distribution/download").status_code == 403

    @pytest.mark.parametrize("who", ["doc", "nurse", "headnurse"])
    def test_publish_is_forbidden(self, client, who):
        login(client, who)
        assert client.post("/distribution/publish").status_code == 403

    @pytest.mark.parametrize("who", ["doc", "nurse", "headnurse"])
    def test_menu_item_is_hidden(self, client, who):
        """Пункта меню нет ни у врача, ни у медсестры, ни у старшей сестры."""
        login(client, who)
        assert 'href="/distribution"' not in client.get("/").text


class TestApi:
    def test_requires_login(self, client):
        assert client.get("/distribution/api/aggregate").status_code == 401

    def test_head_gets_everyone(self, client):
        login(client, "head")
        payload = client.get("/distribution/api/aggregate").json()
        assert payload["by"] == "doctor"
        assert payload["total"] == 0                      # записей ещё никто не подал

    def test_unknown_split_is_400(self, client):
        login(client, "head")
        response = client.get("/distribution/api/aggregate?by=base")
        assert response.status_code == 400
        assert "by" in response.json()["detail"]

    def test_period_from_parameters(self, client):
        login(client, "head")
        payload = client.get("/distribution/api/aggregate?from=2020-01-01&to=2020-12-31").json()
        assert payload["from"] == "2020-01-01" and payload["to"] == "2020-12-31"


class TestSeesRecordsFromNeighbourModule:
    def test_record_added_on_main_page_appears_in_distribution(self, client):
        """Сквозная проверка шва: запись из модуля `records` видна «Распределению»."""
        login(client, "doc")
        created = client.post(
            "/anesthesia",
            data={"patient_name": "Пациент А.", "nurse_id": str(client.ids["nurse"].id)},
        )
        assert created.status_code == 303
        client.post("/logout")

        login(client, "head")
        payload = client.get("/distribution/api/aggregate").json()
        assert payload["total"] == 1
        assert payload["rows"] == [
            {"key": str(client.ids["doctor"].id), "count": 1, "name": "Петров Пётр"}
        ]


class TestVedomost:
    """Прогон, скачивание и публикация — три отдельных действия (ADR-0027)."""

    def _post(self, client, month: str, payload: bytes, name: str = "ведомость.xlsx"):
        return client.post(
            "/distribution/spread",
            data={"month": month},
            files={"file": (name, payload, XLSX_TYPE)},
        )

    def _month(self) -> str:
        return date.today().strftime("%Y-%m")

    def _source(self, *patients: str) -> bytes:
        """Ведомость по строке на пациента — те же суммы, что в настоящей."""
        return vedomost_xlsx([
            {"date": date.today(), "patient": patient,
             "doctor": 1206.8, "smp": 431, "mmp": 86.2}
            for patient in patients
        ])

    def _submit_as_doctor(self, client, patient: str) -> None:
        login(client, "doc")
        client.post(
            "/anesthesia",
            data={"patient_name": patient, "nurse_id": str(client.ids["nurse"].id)},
        )
        client.post("/logout")

    def _doctor_page(self, client) -> str:
        """Страница врача: по ней видно, стоит ли у записи метка «учтено»."""
        client.post("/logout")
        login(client, "doc")
        text = client.get("/").text
        client.post("/logout")
        return text

    def test_anonymous_is_401(self, client):
        assert self._post(client, "2026-06", b"x").status_code == 401

    def test_run_shows_the_result_and_gives_no_file(self, client):
        """Прогон показывает результат на странице и файла не отдаёт."""
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")

        login(client, "head")
        response = self._post(client, self._month(), self._source("ПЕТРОВ П.С."))

        assert response.status_code == 200
        assert "spreadsheet" not in response.headers.get("content-type", "")
        assert "Строк в ведомости" in response.text
        assert "Петров Пётр" in response.text          # сумма по врачу из записи
        assert "Сидорова Анна" in response.text        # и по сестре
        assert "Скачать файл" in response.text
        assert "Опубликовать</button>" in response.text

    def test_run_alone_does_not_mark(self, client):
        """Месяц закрывает публикация, а не загрузка файла (ADR-0027)."""
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")
        login(client, "head")
        assert self._post(client, self._month(), self._source("ПЕТРОВ П.С.")).status_code == 200

        assert "учтено" not in self._doctor_page(client)

    def test_publish_marks_the_record_for_both_staff(self, client):
        """Публикация ставит «учтено» — врач и медсестра это видят."""
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")
        login(client, "head")
        self._post(client, self._month(), self._source("ПЕТРОВ П.С."))

        published = client.post("/distribution/publish")

        assert published.status_code == 200
        assert "помечено записей — 1" in published.text
        assert "учтено" in self._doctor_page(client)

        login(client, "nurse")
        assert "учтено" in client.get("/").text

    def test_download_gives_the_file_with_two_sheets(self, client):
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")
        login(client, "head")
        self._post(client, self._month(), self._source("ПЕТРОВ П.С."))

        response = client.get("/distribution/download")

        assert response.status_code == 200
        assert "spreadsheet" in response.headers["content-type"]
        disposition = unquote(response.headers["content-disposition"])
        assert "распределение" in disposition
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["Sheet Name Here", "Распределение", "Суммы по лицам"]
        row = workbook["Распределение"][2]
        assert row[4].value == "Петров Пётр"           # врач из поданной записи
        assert row[5].value == "Сидорова Анна"         # сестра из неё же
        assert float(row[7].value) == pytest.approx(517.2)

    def test_rerun_alone_keeps_the_published_mark(self, client):
        """Повторный прогон месяца сам по себе метку не снимает (ADR-0027)."""
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")
        month = self._month()
        login(client, "head")
        assert self._post(client, month, self._source("ПЕТРОВ П.С.")).status_code == 200
        assert client.post("/distribution/publish").status_code == 200

        # второй прогон: нашего пациента в ведомости больше нет, публикации нет
        assert self._post(client, month, self._source("МИХАЙЛОВА А.И.")).status_code == 200

        assert "учтено" in self._doctor_page(client)

    def test_publish_after_the_second_run_clears_the_mark(self, client):
        """А публикация второго прогона метку снимает: прежние метки месяца снимаются."""
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")
        month = self._month()
        login(client, "head")
        self._post(client, month, self._source("ПЕТРОВ П.С."))
        client.post("/distribution/publish")

        self._post(client, month, self._source("МИХАЙЛОВА А.И."))
        assert client.post("/distribution/publish").status_code == 200

        assert "учтено" not in self._doctor_page(client)

    def test_attention_row_is_shown_and_gets_no_mark(self, client):
        """Фамилия сошлась, инициалы нет: строка «Внимание», метки нет."""
        self._submit_as_doctor(client, "Петров Пётр Сергеевич")

        login(client, "head")
        response = self._post(client, self._month(), self._source("ПЕТРОВ П.И."))

        assert response.status_code == 200
        assert "Строки «Внимание»" in response.text
        assert "ПЕТРОВ П.И." in response.text
        assert client.post("/distribution/publish").status_code == 200

        assert "учтено" not in self._doctor_page(client)

    def test_page_shows_the_last_run(self, client):
        """Результат последнего прогона виден и на странице, а не только в ответе."""
        login(client, "head")
        self._post(client, self._month(), self._source("МИХАЙЛОВА А.И."))

        text = client.get("/distribution").text

        assert "Результат прогона за" in text
        assert "без пары" in text
        assert "ведомость.xlsx" in text                 # имя загруженного файла

    def test_run_without_records_explains_itself(self, client):
        """Ноль поданных анестезий — на странице сказано, что сверять не с чем."""
        login(client, "head")
        response = self._post(client, self._month(), self._source("МИХАЙЛОВА А.И."))

        assert "поданных анестезий за месяц: <b>0</b>" in response.text
        assert "нет ни одной поданной анестезии" in response.text

    def test_publish_without_a_run_is_a_message(self, client):
        login(client, "head")
        response = client.post("/distribution/publish")
        assert response.status_code == 400
        assert "Прогона нет" in response.text

    def test_download_without_a_run_is_a_message(self, client):
        login(client, "head")
        response = client.get("/distribution/download")
        assert response.status_code == 400
        assert "Прогона нет" in response.text

    def test_foreign_file_is_a_message_not_a_crash(self, client):
        login(client, "head")
        response = self._post(client, "2026-06", "не таблица".encode("utf-8"))
        assert response.status_code == 400
        assert "прочитать" in response.text

    def test_bad_month_is_a_message(self, client):
        login(client, "head")
        source = self._source("ПЕТРОВ П.С.")
        response = self._post(client, "июнь", source)
        assert response.status_code == 400
        assert "месяц" in response.text.lower()

    def test_failed_run_keeps_the_previous_result(self, client):
        """Неудачная загрузка не стирает прежний прогон: он остаётся на странице."""
        login(client, "head")
        self._post(client, self._month(), self._source("МИХАЙЛОВА А.И."))

        response = self._post(client, "июнь", self._source("ПЕТРОВ П.С."))

        assert response.status_code == 400
        assert "Результат прогона за" in response.text
