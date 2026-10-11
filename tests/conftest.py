import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

import os

import pytest


@pytest.fixture(autouse=True, scope="session")
def _isolated_module_dirs(tmp_path_factory):
    """Тесты не пишут в `data/` репозитория.

    При сборке приложения модуль «Чат» создаёт папку PDF-источников
    (ADR-0020) — по умолчанию это `data/wiki/sources` внутри проекта, а
    «Распределение» кладёт туда же прогон ведомости (ADR-0027, `data/distribution`).
    В тестах уводим обе во временную папку: репозиторий остаётся чистым.
    """
    base = tmp_path_factory.mktemp("module-dirs")
    previous = {
        name: os.environ.get(name)
        for name in ("WIKI_SOURCES_DIR", "DISTRIBUTION_DIR")
    }
    os.environ["WIKI_SOURCES_DIR"] = str(base / "wiki-sources")
    os.environ["DISTRIBUTION_DIR"] = str(base / "distribution")
    yield
    for name, value in previous.items():
        if value is None:
            os.environ.pop(name, None)
        else:
            os.environ[name] = value


@pytest.fixture(autouse=True)
def _run_dir_per_test(tmp_path, monkeypatch):
    """Прогон ведомости — у каждого теста свой каталог (ADR-0027).

    Прогон один на раздел и переживает перезапуск приложения: с общим каталогом
    тесты видели бы прогон соседнего теста, и «прогона нет» стало бы
    непроверяемым — именно так это и всплыло при первом же прогоне.
    """
    monkeypatch.setenv("DISTRIBUTION_DIR", str(tmp_path / "distribution"))


@pytest.fixture(scope="session")
def sample_docx(tmp_path_factory):
    from docx import Document
    p = tmp_path_factory.mktemp("docs") / "prikaz_123.docx"
    doc = Document()
    doc.add_heading("ПРИКАЗ № 123 от 15.03.2024", level=0)
    doc.add_paragraph("Об утверждении порядка работы отделения анестезиологии.")
    doc.add_heading("1. Общие положения", level=1)
    doc.add_paragraph("1.1. Настоящий приказ регулирует порядок оказания анестезиологической помощи.")
    doc.add_paragraph("1.2. Дежурный анестезиолог обязан осмотреть пациента до операции и зафиксировать данные в карте.")
    doc.save(p)
    return p


@pytest.fixture(scope="session")
def sample_pdf(tmp_path_factory):
    import pymupdf
    p = tmp_path_factory.mktemp("docs") / "prikaz_456.pdf"
    font = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "ПРИКАЗ № 456 от 20.05.2024", fontsize=16, fontname="arial-uni", fontfile=font)
    page.insert_text((72, 100), "О графике дежурств отделения анестезиологии.", fontsize=11, fontname="arial-uni", fontfile=font)
    page.insert_text((72, 128), "1.1. График дежурств утверждается заведующим отделением ежемесячно.", fontsize=11, fontname="arial-uni", fontfile=font)
    doc.save(p)
    doc.close()
    return p
