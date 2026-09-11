"""Тесты аналитики по истории заявок (analytics.py): слияние недель, фильтры.

Фикстура как в test_needs_report.py: временный YAML-каталог (2 базы × 2
точки; группы Неспецифика, Медикаменты, Растворы) + SQLite-БД на tmp_path
+ NeedsService — заявки сохраняются со снимками unit/grp как в проде
(ТЗ: отчёты по снимкам). Недели две: '2026-08-10' и '2026-08-17'.
Заявки раздельные по разделу (category).
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from docapp.domain.employee import NURSE
from docapp.needs.analytics import period_label, summarize
from docapp.needs.catalog import (
    CATEGORY_MEDICAMENTS,
    CATEGORY_SOLUTIONS,
    SOLUTIONS_GROUP,
)
from docapp.needs.report import build_xlsx
from docapp.needs.service import NeedsService
from docapp.needs.store import SqliteNeedsStore
from factories import catalog, make_db

CATALOG_YAML = """\
bases:
  Ленская: [травма, урология]
  Таймырская: [экстренная, гной]

groups:
  Неспецифика:
    Атропин: амп
    Дексаметазон: амп
  Медикаменты:
    Адреналин: амп
  Растворы:
    Физ 200/250: фл
    Рингер: фл
"""

WEEK1 = "2026-08-10"
WEEK2 = "2026-08-17"

SOL = CATEGORY_SOLUTIONS
MED = CATEGORY_MEDICAMENTS


@pytest.fixture
def env(tmp_path):
    """Каталог + хранилище + сервис на временных файлах (2 базы × 2 точки)."""
    db = make_db(tmp_path)
    catalog_obj = catalog(db, CATALOG_YAML)
    store = SqliteNeedsStore(db)
    service = NeedsService(store, catalog_obj)
    yield service, store, catalog_obj
    store.close_conn()


def send(service, base, point, week, category, lines):
    """Сохранить заявку раздела медсестрой и отправить её (status='sent')."""
    service.save(
        user_id=7,
        role=NURSE,
        base=base,
        point=point,
        category=category,
        week_start=week,
        lines=[{"item": item, "qty": qty} for item, qty in lines],
    )
    service.submit(7, NURSE, base, point, category, week)


class TestSummarize:
    def test_merges_two_weeks(self, env):
        """a) Слияние двух недель: препарат суммируется; растворы поточково с ИТОГО."""
        service, store, catalog = env
        send(service, "Ленская", "травма", WEEK1, MED, [("Атропин", 10)])
        send(service, "Ленская", "травма", WEEK1, SOL, [("Физ 200/250", 5)])
        send(service, "Ленская", "урология", WEEK1, SOL, [("Рингер", 4)])
        send(service, "Ленская", "травма", WEEK2, MED, [("Атропин", 3)])
        send(service, "Ленская", "травма", WEEK2, SOL, [("Физ 200/250", 2)])
        summ = summarize(store.list_range(WEEK1, WEEK2), catalog, WEEK1, WEEK2)

        # остальные группы — суммарно за период
        assert summ["groups"]["Неспецифика"]["Атропин"] == {"unit": "амп", "qty": 13}
        # растворы — поточково за период, у отсутствующей точки 0, ИТОГО за период
        assert summ["solutions"]["Физ 200/250"] == {
            "травма": 7, "урология": 0, "ИТОГО": 7, "unit": "фл",
        }
        assert summ["solutions"]["Рингер"] == {
            "травма": 0, "урология": 4, "ИТОГО": 4, "unit": "фл",
        }

    def test_points_all_touched_bases(self, env):
        """b) 'points' — точки всех затронутых баз в порядке каталога."""
        service, store, catalog = env
        send(service, "Ленская", "травма", WEEK1, MED, [("Атропин", 1)])
        send(service, "Таймырская", "экстренная", WEEK1, MED, [("Атропин", 2)])
        summ = summarize(store.list_range(WEEK1, WEEK2), catalog, WEEK1, WEEK2)

        assert summ["points"] == ["травма", "урология", "экстренная", "гной"]

    def test_base_filter_only_lenskaya(self, env):
        """c) base='Ленская': точки только Ленской, заявки Таймырской не учитываются."""
        service, store, catalog = env
        send(service, "Ленская", "травма", WEEK1, MED, [("Атропин", 10)])
        send(service, "Таймырская", "экстренная", WEEK1, MED, [("Адреналин", 5)])
        # роутер отфильтрует по базе сам — в summarize передаём уже отфильтрованный список
        requests = store.list_range(WEEK1, WEEK2, base="Ленская")
        summ = summarize(requests, catalog, WEEK1, WEEK2, base="Ленская")

        assert summ["points"] == ["травма", "урология"]
        assert "Медикаменты" not in summ["groups"]  # Адреналин Таймырской не учтён
        assert summ["groups"]["Неспецифика"]["Атропин"]["qty"] == 10

    def test_group_filter(self, env):
        """d) group='Растворы' → только solutions; group='Неспецифика' → только она."""
        service, store, catalog = env
        send(service, "Ленская", "травма", WEEK1, MED, [("Атропин", 10)])
        send(service, "Ленская", "травма", WEEK1, SOL, [("Физ 200/250", 5)])
        requests = store.list_range(WEEK1, WEEK2)

        only_solutions = summarize(requests, catalog, WEEK1, WEEK2, group=SOLUTIONS_GROUP)
        assert only_solutions["groups"] == {}
        assert "Физ 200/250" in only_solutions["solutions"]

        only_group = summarize(requests, catalog, WEEK1, WEEK2, group="Неспецифика")
        assert only_group["solutions"] == {}
        assert list(only_group["groups"]) == ["Неспецифика"]
        assert only_group["groups"]["Неспецифика"]["Атропин"]["qty"] == 10

    def test_section_filter(self, env):
        """d') section='solutions' → только растворы; 'medicaments' → только группы."""
        service, store, catalog = env
        send(service, "Ленская", "травма", WEEK1, MED, [("Атропин", 10)])
        send(service, "Ленская", "травма", WEEK1, SOL, [("Физ 200/250", 5)])
        requests = store.list_range(WEEK1, WEEK2)

        only_solutions = summarize(requests, catalog, WEEK1, WEEK2, section=SOL)
        assert only_solutions["section"] == SOL
        assert only_solutions["groups"] == {}
        assert "Физ 200/250" in only_solutions["solutions"]

        only_meds = summarize(requests, catalog, WEEK1, WEEK2, section=MED)
        assert only_meds["section"] == MED
        assert only_meds["solutions"] == {}
        assert only_meds["groups"]["Неспецифика"]["Атропин"]["qty"] == 10

        both = summarize(requests, catalog, WEEK1, WEEK2, section=None)
        assert both["section"] is None
        assert "Физ 200/250" in both["solutions"]
        assert "Атропин" in both["groups"]["Неспецифика"]

    def test_drafts_excluded(self, env):
        """e) Черновики (draft) в аналитику не попадают (их отсекает агрегация)."""
        service, store, catalog = env
        # черновик с реальными количествами — не отправлен
        service.save(8, NURSE, "Ленская", "травма", MED, WEEK1, [{"item": "Атропин", "qty": 10}])
        summ = summarize(store.list_range(WEEK1, WEEK2), catalog, WEEK1, WEEK2)

        assert summ["groups"] == {}
        assert summ["solutions"] == {}

    def test_empty_period(self, env):
        """g) Пустой период: пустые solutions/groups, точки всех баз каталога."""
        service, store, catalog = env
        summ = summarize(store.list_range(WEEK1, WEEK2), catalog, WEEK1, WEEK2)

        assert summ["solutions"] == {}
        assert summ["groups"] == {}
        assert summ["points"] == ["травма", "урология", "экстренная", "гной"]
        assert summ["period_label"] == "10.08.2026 – 17.08.2026"
        assert summ["from_week"] == WEEK1 and summ["to_week"] == WEEK2
        assert summ["base"] is None and summ["group"] is None
        assert summ["section"] is None

    def test_build_xlsx_reused(self, env):
        """h) build_xlsx из report.py работает с агрегатом summarize (листы на месте)."""
        service, store, catalog = env
        send(service, "Ленская", "травма", WEEK1, MED, [("Атропин", 10)])
        send(service, "Ленская", "травма", WEEK1, SOL, [("Физ 200/250", 5)])
        send(service, "Ленская", "урология", WEEK2, SOL, [("Рингер", 4)])
        summ = summarize(store.list_range(WEEK1, WEEK2), catalog, WEEK1, WEEK2)

        wb = load_workbook(BytesIO(build_xlsx(summ)))
        assert wb.sheetnames == ["Растворы", "Медикаменты"]
        sol = wb["Растворы"]
        rows = [[cell.value for cell in row] for row in sol.iter_rows()]
        # данные периода доехали до xlsx: строка раствора с ИТОГО за период
        row_fiz = next(row for row in rows if row and row[0] == "Физ 200/250")
        assert row_fiz[1] == "фл"
        assert row_fiz[-1] == 5


class TestPeriodLabel:
    def test_two_weeks(self):
        """f1) Разные недели: '10.08.2026 – 24.08.2026' (понедельники недель)."""
        assert period_label("2026-08-10", "2026-08-24") == "10.08.2026 – 24.08.2026"

    def test_single_week(self):
        """f2) Одна неделя: одна дата '10.08.2026'."""
        assert period_label("2026-08-10", "2026-08-10") == "10.08.2026"
