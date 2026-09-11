"""Тесты ведомости: разбор файла, разноска по анестезиям и выгрузка (.xlsx).

Файл для тестов синтетический, но повторяет структуру настоящей выгрузки
больницы (`factories.vedomost_xlsx`, ТЗ `docs/ТЗ-распределение.md`): ФИО
пациентов вымышленные, файл владельца в репозиторий не попадает.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from docapp.distribution import xlsx
from docapp.distribution.service import DistributionService
from docapp.distribution.spread import person_key, spread
from docapp.distribution.vedomost import VedomostError, parse
from docapp.domain.anesthesia import Anesthesia
from docapp.domain.employee import DOCTOR, NURSE, Employee
from docapp.people.store import SqliteEmployeeStore
from docapp.records.service import AnesthesiaService
from docapp.records.store import SqliteAnesthesiaStore
from factories import make_db, vedomost_xlsx

#: Ведомость: два ряда одного пациента (осмотр + анестезия), ненайденный
#: пациент, неоднозначная фамилия и майский пациент в июньском файле.
ROWS = [
    {"date": date(2026, 6, 2), "patient": "ПЕТРОВ П.С.",
     "service": "Тотальная внутривенная анестезия(30 мин.)",
     "doctor": 1206.8, "smp": 431, "mmp": 86.2},
    {"date": date(2026, 6, 2), "patient": "ПЕТРОВ П.С.",
     "service": "Осмотр (консультация) врачом-анестезиологом первичный",
     "doctor": 675.81, "smp": 241.36, "mmp": 48.27},
    {"date": date(2026, 6, 3), "patient": "МИХАЙЛОВА А.И.",
     "doctor": 1206.8, "smp": 431, "mmp": 86.2},
    {"date": date(2026, 6, 4), "patient": "КУЗНЕЦОВ П.И.",
     "doctor": 1206.8, "smp": 431, "mmp": 86.2},
    {"date": date(2026, 6, 5), "patient": "ОРЛОВ О.О.",
     "doctor": 1206.8, "smp": 431, "mmp": 86.2},
    {"date": date(2026, 6, 6), "patient": "ГАВРИЛОВ И.П.",
     "doctor": 1206.8, "smp": 431, "mmp": 86.2},
]


def make_anesthesia(**overrides) -> Anesthesia:
    base: dict = dict(
        date=date(2026, 6, 1),
        patient_name="Пациент",
        doctor_id=1,
        nurse_id=2,
        created_at=datetime(2026, 6, 1, 9, 0, tzinfo=timezone.utc),
    )
    base.update(overrides)
    return Anesthesia(**base)


@pytest.fixture
def env(tmp_path):
    """Справочник, поданные анестезии и разобранная ведомость."""
    db = make_db(tmp_path, seed=False)
    employees = SqliteEmployeeStore(db)
    doctor = employees.add(Employee(last_name="Зайцев", first_name="Захар",
                                    middle_name="Захарович", role=DOCTOR, buh_id="D-1"))
    nurse = employees.add(Employee(last_name="Ковалёва", first_name="Анна",
                                   middle_name="Ивановна", role=NURSE, buh_id="S-1"))
    nurse_no_id = employees.add(Employee(last_name="Романова", first_name="Ольга",
                                         middle_name="Сергеевна", role=NURSE))
    records = [
        make_anesthesia(date=date(2026, 6, 2), patient_name="Петров Пётр Сергеевич",
                        doctor_id=doctor.id, nurse_id=nurse.id),
        # три записи с одной фамилией: две с одинаковыми инициалами (неоднозначно)
        # и одна с другими (её можно выбрать по инициалам)
        make_anesthesia(date=date(2026, 6, 3), patient_name="Кузнецов Пётр Иванович",
                        doctor_id=doctor.id, nurse_id=nurse_no_id.id),
        make_anesthesia(date=date(2026, 6, 4), patient_name="Кузнецов Павел Сергеевич",
                        doctor_id=doctor.id, nurse_id=nurse_no_id.id),
        make_anesthesia(date=date(2026, 6, 4), patient_name="Кузнецов Павел Иванович",
                        doctor_id=doctor.id, nurse_id=nurse_no_id.id),
        make_anesthesia(date=date(2026, 6, 6), patient_name="Гаврилов Игорь Петрович",
                        doctor_id=doctor.id, nurse_id=nurse_no_id.id),
        # майская запись: в июньскую ведомость она попасть не должна
        make_anesthesia(date=date(2026, 5, 20), patient_name="Орлов Олег Олегович",
                        doctor_id=doctor.id, nurse_id=nurse.id),
    ]
    yield {"employees": employees, "records": records, "doctor": doctor,
           "nurse": nurse, "nurse_no_id": nurse_no_id}
    employees.close()


def spread_file(env, data: bytes):
    """Разнести синтетическую ведомость по записям июня 2026.

    Сервис берёт записи у модуля `records` строго за месяц (`spread` в сервисе),
    а сама функция разноски работает с уже выбранным набором — поэтому здесь
    записи июня отбираются так же, как это делает интерфейс записей.
    """
    june = [record for record in env["records"] if record.date.month == 6]
    return spread("2026-06", parse(data), june, env["employees"])


# ── месяц — строгая граница (сервис, через интерфейс записей) ─────────


class TestServiceMonth:
    """Сервис берёт записи у модуля `records` строго за выбранный месяц."""

    @pytest.fixture
    def service(self, tmp_path):
        db = make_db(tmp_path, seed=False)
        employees = SqliteEmployeeStore(db)
        doctor = employees.add(Employee(last_name="Зайцев", first_name="Захар", role=DOCTOR))
        nurse = employees.add(Employee(last_name="Ковалёва", first_name="Анна", role=NURSE))
        store = SqliteAnesthesiaStore(db)
        store.add(make_anesthesia(date=date(2026, 5, 20), patient_name="Орлов Олег Олегович",
                                  doctor_id=doctor.id, nurse_id=nurse.id))
        yield DistributionService(AnesthesiaService(store), employees)
        store.close()
        employees.close()

    def _file(self):
        return vedomost_xlsx([{"date": date(2026, 6, 5), "patient": "ОРЛОВ О.О.",
                               "doctor": 1206.8, "smp": 431, "mmp": 86.2}])

    def test_other_month_record_is_not_used(self, service):
        result = service.spread("2026-06", self._file())
        assert result.assigned_rows == 0
        assert result.rows[0].doctor is None

    def test_same_patient_found_in_his_month(self, service):
        result = service.spread("2026-05", self._file())
        assert result.assigned_rows == 1
        assert result.rows[0].doctor.full_name == "Зайцев Захар"
        assert result.doctors[0].amount == Decimal("1206.8")
        assert result.nurses[0].amount == Decimal("517.2")

    def test_bad_month_is_a_value_error(self, service):
        with pytest.raises(ValueError, match="месяц"):
            service.spread("июнь", self._file())


# ── разбор файла ──────────────────────────────────────────────────────


class TestParse:
    def test_rows_dates_and_amounts(self):
        vedomost = parse(vedomost_xlsx(ROWS))
        assert len(vedomost.rows) == len(ROWS)
        assert vedomost.sheet_name == "Sheet Name Here"
        assert vedomost.period == "Период с 01.06.2026 по 30.06.2026"
        first = vedomost.rows[0]
        assert first.service_date == date(2026, 6, 2)
        assert first.patient == "ПЕТРОВ П.С."
        assert first.doctor_amount == Decimal("1206.8")
        # сестре идут «СМП» и «ММП» одной суммой
        assert first.nurse_amount == Decimal("517.2")

    def test_total_row_is_not_a_service(self):
        """«Итого по отделению» — не строка услуги: пациента в ней нет."""
        vedomost = parse(vedomost_xlsx(ROWS))
        assert all("Итого" not in row.patient for row in vedomost.rows)

    def test_sum_given_as_text(self):
        """Суммы могут прийти строкой с запятой — читаем как число."""
        vedomost = parse(vedomost_xlsx([{**ROWS[0], "doctor": "1 206,80"}]))
        assert vedomost.rows[0].doctor_amount == Decimal("1206.80")

    def test_not_a_vedomost(self):
        """Не xlsx — понятная ошибка чтения."""
        with pytest.raises(VedomostError, match="прочитать"):
            parse(b"")
        with pytest.raises(VedomostError, match="прочитать"):
            parse("\x00\x01\x02 — не xlsx".encode("utf-8"))

    def test_xlsx_without_header_is_rejected(self):
        """Таблица без шапки ведомости — не ведомость: так и говорим."""
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.worksheets[0].cell(row=1, column=1, value="просто таблица")
        out = BytesIO()
        workbook.save(out)
        with pytest.raises(VedomostError, match="шапки"):
            parse(out.getvalue())

    def test_non_numeric_amount_names_the_row(self):
        with pytest.raises(VedomostError, match="строка 8"):
            parse(vedomost_xlsx([{**ROWS[0], "doctor": "много"}]))

    def test_no_rows_with_patients(self):
        with pytest.raises(VedomostError, match="ни одной строки"):
            parse(vedomost_xlsx([]))


# ── сопоставление по пациенту ─────────────────────────────────────────


class TestPersonKey:
    @pytest.mark.parametrize(
        "fio,expected",
        [
            ("ПЕТРОВ П.С.", ("петров", "пс")),
            ("Петров Пётр Сергеевич", ("петров", "пс")),
            ("  ПЕТРОВА  А. И. ", ("петрова", "аи")),
            ("", ("", "")),
        ],
    )
    def test_key(self, fio, expected):
        assert person_key(fio) == expected


class TestSpread:
    def test_pair_goes_to_every_row_of_the_patient(self, env):
        """Одна заявка врача закрывает и осмотр, и анестезию (ТЗ §сопоставление)."""
        result = spread_file(env, vedomost_xlsx(ROWS))
        petrov_rows = [row for row in result.rows if row.patient == "ПЕТРОВ П.С."]
        assert len(petrov_rows) == 2
        assert all(row.doctor == env["doctor"] and row.nurse == env["nurse"]
                   for row in petrov_rows)

    def test_unknown_patient_stays_empty(self, env):
        result = spread_file(env, vedomost_xlsx(ROWS))
        mikhailova = next(row for row in result.rows if row.patient == "МИХАЙЛОВА А.И.")
        assert mikhailova.doctor is None and mikhailova.nurse is None
        assert mikhailova.attention is False

    def test_same_surname_marks_attention_without_pair(self, env):
        """Однофамильцы — «Внимание», пары нет: угадывать нельзя."""
        result = spread_file(env, vedomost_xlsx(ROWS))
        kuznetsov = next(row for row in result.rows if row.patient == "КУЗНЕЦОВ П.И.")
        assert kuznetsov.attention is True
        assert kuznetsov.doctor is None and kuznetsov.nurse is None

    def test_record_of_another_month_is_not_used(self, env):
        """Месяц строгий: майская запись в июньской ведомости не находится."""
        result = spread_file(env, vedomost_xlsx(ROWS))
        orlov = next(row for row in result.rows if row.patient == "ОРЛОВ О.О.")
        assert orlov.doctor is None and orlov.attention is False

    def test_initials_disambiguate(self, env):
        """Инициалы уточняют выбор, когда фамилия в месяце не одна."""
        rows = [{"date": date(2026, 6, 4), "patient": "КУЗНЕЦОВ П.С.",
                 "doctor": 100, "smp": 50, "mmp": 10}]
        vedomost = parse(vedomost_xlsx(rows))
        june = [record for record in env["records"] if record.date.month == 6]
        result = spread("2026-06", vedomost, june, env["employees"])
        assert result.rows[0].attention is False
        assert result.rows[0].doctor == env["doctor"]

    def test_totals_by_doctor_and_nurse(self, env):
        """Суммы: врачу — «Врач», сестре — «СМП» + «ММП», по убыванию."""
        result = spread_file(env, vedomost_xlsx(ROWS))
        assert [(person.employee.full_name, person.amount) for person in result.doctors] == [
            ("Зайцев Захар Захарович", Decimal("1206.8") * 2 + Decimal("675.81")),
        ]
        assert [(person.employee.full_name, person.amount) for person in result.nurses] == [
            ("Ковалёва Анна Ивановна", Decimal("517.2") + Decimal("289.63")),
            ("Романова Ольга Сергеевна", Decimal("517.2")),
        ]

    def test_counters(self, env):
        result = spread_file(env, vedomost_xlsx(ROWS))
        assert len(result.rows) == 6
        assert result.assigned_rows == 3
        assert result.attention_rows == 1
        assert result.unmatched_rows == 2

    def test_people_without_buh_id_are_listed(self, env):
        result = spread_file(env, vedomost_xlsx(ROWS))
        assert [employee.full_name for employee in result.without_buh_id()] == [
            "Романова Ольга Сергеевна"
        ]


# ── выгрузка ──────────────────────────────────────────────────────────


class TestXlsx:
    @pytest.fixture
    def result(self, env):
        source = vedomost_xlsx(ROWS)
        spread_result = spread_file(env, source)
        return source, xlsx.build(source, spread_result), spread_result

    def test_source_sheets_are_kept(self, result):
        source, built, _ = result
        assert load_workbook(BytesIO(source)).sheetnames == ["Sheet Name Here"]
        assert load_workbook(BytesIO(built)).sheetnames == [
            "Sheet Name Here", xlsx.ROWS_SHEET, xlsx.TOTALS_SHEET
        ]

    def test_rows_sheet_has_pairs_and_sums(self, result):
        _, built, _ = result
        sheet = load_workbook(BytesIO(built))[xlsx.ROWS_SHEET]
        assert [cell.value for cell in sheet[1]] == list(xlsx.ROW_HEADERS)
        first = sheet[2]
        assert first[0].value == 8                     # номер строки ведомости
        assert first[1].value == "02.06.2026"
        assert first[2].value == "ПЕТРОВ П.С."
        assert first[4].value == "Зайцев Захар Захарович"
        assert first[5].value == "Ковалёва Анна Ивановна"
        assert float(first[6].value) == pytest.approx(1206.8)
        assert float(first[7].value) == pytest.approx(517.2)

    def test_attention_cell_is_highlighted(self, result):
        _, built, _ = result
        sheet = load_workbook(BytesIO(built))[xlsx.ROWS_SHEET]
        attention = [
            row for row in sheet.iter_rows(min_row=2)
            if row[8].value == xlsx.ATTENTION
        ]
        assert len(attention) == 1
        cell = attention[0][8]
        assert cell.font.bold is True
        assert cell.fill.start_color.rgb.endswith("FFF3CD")
        assert attention[0][6].value is None            # суммы не проставлены

    def test_totals_sheet_lists_people(self, result):
        _, built, _ = result
        sheet = load_workbook(BytesIO(built))[xlsx.TOTALS_SHEET]
        text = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
        assert "Врачи" in text and "Сёстры" in text
        assert "Зайцев Захар Захарович" in text
        assert "Ковалёва Анна Ивановна" in text
        assert any("Без ID в бухгалтерии" in value for value in text)
        assert "Романова Ольга Сергеевна" in text
        assert any("Строк в ведомости: 6" in value and "без пары: 2" in value for value in text)

    def test_second_run_does_not_duplicate_sheets(self, env):
        """Файл можно прогнать дважды: наши листы перезаписываются."""
        source = vedomost_xlsx(ROWS)
        first = xlsx.build(source, spread_file(env, source))
        second = xlsx.build(first, spread_file(env, source))
        assert load_workbook(BytesIO(second)).sheetnames == [
            "Sheet Name Here", xlsx.ROWS_SHEET, xlsx.TOTALS_SHEET
        ]
