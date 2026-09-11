"""Разлиновка дежурств: сетка «время × врач» в .xlsx.

Лист на каждую дату смены. Шапка: «Время» + часы 16…8 (каждый час = 4 колонки
по 15 минут — итого 68 колонок, как в шаблоне «Разлиновка.xlsx»). Секции по
базам («Таймырская», «Ленская»), в каждой — строки врачей; операция рисуется
полоской (объединённые ячейки) от времени начала до времени конца.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from docapp.core import period

#: Часы в шапке разлиновки (17 подписей: 16…23, 0…8). «8» — граница 08:00.
HOURS = [16, 17, 18, 19, 20, 21, 22, 23, 0, 1, 2, 3, 4, 5, 6, 7, 8]

#: Колонок на час (15-минутные слоты).
SLOTS_PER_HOUR = 4

#: Первая колонка сетки (колонка A = 1 — имена врачей / «Время» / база).
FIRST_TIME_COL = 2

#: Последняя колонка сетки.
LAST_TIME_COL = FIRST_TIME_COL + len(HOURS) * SLOTS_PER_HOUR - 1

_HEADER_FONT = Font(bold=True)
_BASE_FONT = Font(bold=True)
_OP_FONT = Font(size=9)
_OP_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HEADER_ALIGN = Alignment(horizontal="center")
_THIN = Side(style="thin", color="B0B0B0")
_OP_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BASE_FILL = PatternFill("solid", fgColor="DDEBF7")
_HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")


def _slot_index(hhmm: str) -> int:
    """'HH:MM' → индекс 15-минутного слота от 16:00 (0..64)."""
    return period.minutes_into_shift(hhmm) // 15


def _date_title(iso_date: str) -> str:
    """'YYYY-MM-DD' → 'DD.MM' (имя листа, как в шаблоне)."""
    y, m, d = iso_date.split("-")
    return f"{d}.{m}"


def build_xlsx(days: list[dict]) -> bytes:
    """Собрать .xlsx разлиновки и вернуть bytes.

    days — список в хронологическом порядке, каждый:
        {"date": "YYYY-MM-DD",
         "bases": [{"base": "Таймырская",
                    "doctors": [{"name": "Фамилия И.О.",
                                 "operations": [{"operation", "start_time", "end_time"}]}]}]}
    """
    wb = Workbook()
    active = wb.active
    assert active is not None  # свежий Workbook всегда с активным листом
    wb.remove(active)

    for day in days:
        ws = wb.create_sheet(title=_date_title(day["date"]))
        ws.column_dimensions["A"].width = 14
        for col in range(FIRST_TIME_COL, LAST_TIME_COL + 1):
            ws.column_dimensions[get_column_letter(col)].width = 3

        # Шапка: «Время» + часы (каждый час — 4 объединённые колонки).
        cell = ws.cell(row=1, column=1, value="Время")
        cell.font = _HEADER_FONT
        for i, hour in enumerate(HOURS):
            c1 = FIRST_TIME_COL + i * SLOTS_PER_HOUR
            c2 = c1 + SLOTS_PER_HOUR - 1
            ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
            hcell = ws.cell(row=1, column=c1, value=hour)
            hcell.font = _HEADER_FONT
            hcell.alignment = _HEADER_ALIGN
            hcell.fill = _HEADER_FILL

        row = 2
        for base_entry in day["bases"]:
            # Заголовок базы на всю ширину сетки.
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST_TIME_COL)
            bcell = ws.cell(row=row, column=1, value=base_entry["base"])
            bcell.font = _BASE_FONT
            bcell.fill = _BASE_FILL
            row += 1

            for doctor in base_entry["doctors"]:
                ws.cell(row=row, column=1, value=doctor["name"])
                _place_operations(ws, row, doctor["operations"])
                row += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _place_operations(ws, row: int, operations: list[dict]) -> None:
    """Нарисовать операции врача полосками в строке.

    Пересекающиеся полоски (редко, по ошибке ввода) объединяются в одну —
    названия через запятую — чтобы не получить битых объединённых ячеек.
    """
    ordered = sorted(operations, key=lambda op: _slot_index(op["start_time"]))
    placed: list[list] = []  # [c1, c2, [названия]]
    for op in ordered:
        s = _slot_index(op["start_time"])
        e = _slot_index(op["end_time"])
        c1 = FIRST_TIME_COL + s
        c2 = FIRST_TIME_COL + e - 1  # конец полоски (включительно)
        if c2 < c1:
            c2 = c1  # операция короче слота — одна ячейка
        merged_into = None
        for p in placed:
            if c1 <= p[1] and c2 >= p[0]:  # пересечение по колонкам
                merged_into = p
                break
        if merged_into is not None:
            merged_into[0] = min(merged_into[0], c1)
            merged_into[1] = max(merged_into[1], c2)
            merged_into[2].append(op["operation"])
        else:
            placed.append([c1, c2, [op["operation"]]])

    for c1, c2, names in placed:
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        cell = ws.cell(row=row, column=c1, value=", ".join(names))
        cell.font = _OP_FONT
        cell.alignment = _OP_ALIGN
        cell.border = _OP_BORDER
