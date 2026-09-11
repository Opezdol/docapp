"""Выгрузка распределения: тот же файл плюс два листа.

Исходные листы не трогаем: открываем присланную ведомость как есть и дописываем
«Распределение» (построчно) и «Суммы по лицам» (врачи и сёстры отдельными
таблицами — решение владельца). Если файл прогнали дважды, наши листы
перезаписываются, а не двоятся.

Суммы пишем числами с форматом «#,##0.00»: считаем в `Decimal`, а в книгу
кладём число, потому что Excel хранит числа сам.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from docapp.distribution.spread import Person, Spread, SpreadRow

#: Наши листы: построчная разноска и итоги по людям.
ROWS_SHEET = "Распределение"
TOTALS_SHEET = "Суммы по лицам"
#: Пометка неоднозначной строки: жёлтая заливка и жирный текст.
ATTENTION = "Внимание"

ROW_HEADERS = ("№ строки", "Дата", "Пациент", "Услуга", "Врач", "Сестра",
               "Сумма врача", "Сумма сестры", "Пометка")
PERSON_HEADERS = ("ФИО", "ID в бухгалтерии", "Сумма")
ROW_WIDTHS = (10, 12, 30, 48, 26, 26, 14, 14, 12)
PERSON_WIDTHS = (34, 20, 14)
DATE_FORMAT = "%d.%m.%Y"
MONEY_FORMAT = "#,##0.00"

_HEADER_FONT = Font(bold=True)
_ATTENTION_FONT = Font(bold=True)
_ATTENTION_FILL = PatternFill("solid", fgColor="FFF3CD")


def build(source: bytes, spread: Spread) -> bytes:
    """Собрать файл результата: исходные листы плюс два наших."""
    workbook = load_workbook(BytesIO(source))
    for name in (ROWS_SHEET, TOTALS_SHEET):
        if name in workbook.sheetnames:      # повторный прогон — перезаписываем
            del workbook[name]
    _rows_sheet(workbook.create_sheet(ROWS_SHEET), spread)
    _totals_sheet(workbook.create_sheet(TOTALS_SHEET), spread)
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def filename(month: str) -> str:
    """Имя скачиваемого файла: месяц, за который разнесли ведомость."""
    return f"распределение-{month}.xlsx"


def _widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    """Ширина колонок: иначе длинные названия услуг не читаются."""
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _header(sheet: Worksheet, headers: tuple[str, ...]) -> None:
    """Строка заголовков таблицы."""
    sheet.append(list(headers))
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT


def _date_text(row: SpreadRow) -> str:
    return row.service_date.strftime(DATE_FORMAT) if row.service_date else ""


def _money(sheet: Worksheet, row: int, column: int, amount) -> None:
    """Сумма числом с денежным форматом."""
    cell = sheet.cell(row=row, column=column, value=amount)
    cell.number_format = MONEY_FORMAT


def _rows_sheet(sheet: Worksheet, spread: Spread) -> None:
    """Лист «Распределение»: строка ведомости — врач, сестра и суммы."""
    _header(sheet, ROW_HEADERS)
    for row in spread.rows:
        sheet.append([
            row.source_row,
            _date_text(row),
            row.patient,
            row.service,
            row.doctor.full_name if row.doctor else "",
            row.nurse.full_name if row.nurse else "",
            None,
            None,
            ATTENTION if row.attention else "",
        ])
        line = sheet.max_row
        if row.doctor:
            _money(sheet, line, 7, float(row.doctor_amount))
        if row.nurse:
            _money(sheet, line, 8, float(row.nurse_amount))
        if row.attention:
            cell = sheet.cell(row=line, column=len(ROW_HEADERS))
            cell.font = _ATTENTION_FONT
            cell.fill = _ATTENTION_FILL
    _widths(sheet, ROW_WIDTHS)


def _person_table(sheet: Worksheet, title: str, persons: tuple[Person, ...]) -> None:
    """Таблица итогов по людям: подпись, шапка, строки."""
    sheet.append([title])
    sheet.cell(row=sheet.max_row, column=1).font = _HEADER_FONT
    _header(sheet, PERSON_HEADERS)
    for person in persons:
        sheet.append([
            person.employee.full_name,
            person.employee.buh_id or "",
            None,
        ])
        _money(sheet, sheet.max_row, len(PERSON_HEADERS), float(person.amount))


def _totals_sheet(sheet: Worksheet, spread: Spread) -> None:
    """Лист «Суммы по лицам»: счётчики, врачи, сёстры и предупреждения."""
    sheet.append([f"Распределение за {spread.month}"])
    sheet.cell(row=1, column=1).font = _HEADER_FONT
    sheet.append([
        f"Строк в ведомости: {len(spread.rows)} · с врачом и сестрой: {spread.assigned_rows}"
        f" · «Внимание»: {spread.attention_rows} · без пары: {spread.unmatched_rows}"
    ])
    sheet.append([])
    _person_table(sheet, "Врачи", spread.doctors)
    sheet.append([])
    _person_table(sheet, "Сёстры", spread.nurses)

    without_id = spread.without_buh_id()
    if without_id:
        sheet.append([])
        sheet.append(["Без ID в бухгалтерии — бухгалтерия не примет:"])
        sheet.cell(row=sheet.max_row, column=1).font = _ATTENTION_FONT
        for employee in without_id:
            sheet.append([employee.full_name])
    _widths(sheet, PERSON_WIDTHS)
