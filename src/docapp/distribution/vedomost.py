"""Разбор файла ведомости: строки услуг, суммы и период.

Формат подтверждён на реальной выгрузке больницы (см. `docs/ТЗ-распределение.md`):
шапка в седьмой строке, ниже строки услуг, последняя строка — «Итого по
отделению». Числа могут прийти числами, строками или датами — здесь всё
приводится к типам, а непонятное отвергается понятной ошибкой, а не падением.

Суммы читаются из колонок файла и не пересчитываются: «Врач» — врачу, «СМП» плюс
«ММП» — сестре (решение владельца).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

#: Обязательные колонки: нормализованная подпись → что это.
REQUIRED = {
    "фио пациента": "пациент",
    "название услуги": "услуга",
    "врач": "сумма врачу",
    "смп": "сумма сестре",
    "ммп": "сумма сестре",
}
#: Необязательная, но полезная: дата оказания услуги.
DATE_COLUMN = "дата оказания услуги"
#: Сколько строк сверху просматриваем в поисках шапки.
HEADER_SEARCH_ROWS = 20


class VedomostError(ValueError):
    """Файл не похож на ведомость — понятное сообщение вместо падения."""


@dataclass(frozen=True)
class Row:
    """Строка ведомости: какая услуга, кому и сколько."""

    number: int                     # номер строки в исходном файле — для сверки
    service_date: date | None
    patient: str                    # как записано в файле
    service: str
    doctor_amount: Decimal          # колонка «Врач»
    nurse_amount: Decimal           # «СМП» + «ММП»


@dataclass(frozen=True)
class Vedomost:
    """Разобранная ведомость: строки и то, что о ней известно."""

    rows: tuple[Row, ...]
    sheet_name: str
    period: str                     # строка «Период с … по …», если нашлась


def _label(value: object) -> str:
    """Подпись колонки к сравнимому виду: регистр, ё → е, точки и запятые прочь."""
    text = str(value or "").strip().lower().replace("ё", "е")
    for char in ".,;:":
        text = text.replace(char, "")
    return re.sub(r"\s+", " ", text).strip()


def _money(value: object, column: str, row: int) -> Decimal:
    """Сумма из ячейки; пусто — ноль, непонятное — ошибка со строкой."""
    if value is None or str(value).strip() == "":
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        raise VedomostError(
            f"строка {row}: в столбце «{column}» не число: {value!r}"
        ) from None


def _as_date(value: object) -> date | None:
    """Дата оказания услуги, если её удалось разобрать."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _find_header(sheet) -> tuple[int, dict[str, int]]:
    """Найти шапку и колонки: подпись → номер колонки."""
    for row in range(1, min(HEADER_SEARCH_ROWS, sheet.max_row) + 1):
        columns: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            columns.setdefault(_label(sheet.cell(row=row, column=column).value), column)
        if "фио пациента" in columns:
            return row, columns
    raise VedomostError(
        "в файле нет шапки со столбцом «Ф.И.О. пациента» — это не ведомость "
        "больницы или изменился её формат"
    )


def _period(sheet, header_row: int) -> str:
    """Строка «Период с … по …» над шапкой — как написано, для протокола."""
    for row in range(1, header_row):
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row=row, column=column).value
            if isinstance(value, str) and value.strip().lower().startswith("период"):
                return value.strip()
    return ""


def parse(data: bytes) -> Vedomost:
    """Разобрать ведомость: строки услуг, суммы, период.

    Читается первый лист: у больницы это ведомость, а рабочие листы владельца
    («data») мы не трогаем.
    """
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 — любой сбой чтения = «не тот файл»
        raise VedomostError(f"не удалось прочитать файл как .xlsx: {exc}") from exc

    sheet = workbook.worksheets[0]
    header_row, columns = _find_header(sheet)

    missing = [name for name in REQUIRED if name not in columns]
    if missing:
        raise VedomostError(
            "в ведомости нет обязательных столбцов: " + ", ".join(sorted(missing))
        )

    patient_col = columns["фио пациента"]
    rows: list[Row] = []
    for number in range(header_row + 1, sheet.max_row + 1):
        patient = sheet.cell(row=number, column=patient_col).value
        if patient is None or not str(patient).strip():
            continue
        doctor_amount = _money(sheet.cell(row=number, column=columns["врач"]).value,
                              "Врач", number)
        nurse_amount = _money(sheet.cell(row=number, column=columns["смп"]).value,
                              "СМП", number) + _money(
            sheet.cell(row=number, column=columns["ммп"]).value, "ММП", number)
        date_col = columns.get(DATE_COLUMN)
        rows.append(
            Row(
                number=number,
                service_date=_as_date(sheet.cell(row=number, column=date_col).value)
                if date_col else None,
                patient=str(patient).strip(),
                service=str(sheet.cell(row=number, column=columns["название услуги"]).value
                            or "").strip(),
                doctor_amount=doctor_amount,
                nurse_amount=nurse_amount,
            )
        )

    if not rows:
        raise VedomostError("в ведомости не нашлось ни одной строки с пациентом")

    return Vedomost(rows=tuple(rows), sheet_name=sheet.title, period=_period(sheet, header_row))
