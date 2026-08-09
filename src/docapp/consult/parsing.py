"""Парсинг приказов: .docx и .pdf в список блоков Block."""

from dataclasses import dataclass

from docx import Document


@dataclass
class Block:
    kind: str          # "heading" | "para" | "table"
    text: str
    level: int = 0


def parse_docx(path) -> list[Block]:
    """Прочитать .docx: абзацы (стили Heading* -> heading) и таблицы."""
    doc = Document(path)
    blocks = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        if "heading" in style or "заголовок" in style or "title" in style:
            digits = "".join(ch for ch in style if ch.isdigit())
            blocks.append(Block("heading", text, level=int(digits) if digits else 1))
        else:
            blocks.append(Block("para", text))
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(x for x in cells if x)
            if line:
                blocks.append(Block("table", line))
    return blocks


def parse_pdf(path) -> list[Block]:
    """Прочитать .pdf: заголовки определяются по размеру шрифта (>= 13.5)."""
    import pymupdf
    doc = pymupdf.open(path)
    blocks = []
    for page in doc:
        for raw in page.get_text("dict")["blocks"]:
            if raw.get("type") != 0:
                continue
            for line in raw["lines"]:
                spans = line["spans"]
                if not spans:
                    continue
                text = "".join(s["text"] for s in spans).strip()
                if not text:
                    continue
                size = max(s["size"] for s in spans)
                if size >= 13.5:
                    blocks.append(Block("heading", text, level=1))
                else:
                    blocks.append(Block("para", text))
    doc.close()
    return blocks
